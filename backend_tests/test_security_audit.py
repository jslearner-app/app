# Iteration 10: Security audit fixes - CORS allowlist, CSRF, forgot/reset password, formula injection,
# payments/status auth, trusted origin_url, rate limits
import csv
import glob
import io
import uuid

import pytest
import requests
from openpyxl import load_workbook
from test_ledgerstream_regression import API, BASE_URL, DEMO, mdb

EVIL = "https://evil.example"
FORMULA = '=HYPERLINK("x")'


def register(email=None, password="TestPass#2026"):
    s = requests.Session()
    email = email or f"test_sec_{uuid.uuid4().hex[:8]}@example.com"
    r = s.post(f"{API}/auth/register", json={"email": email, "password": password, "name": "TEST Sec", "firm_name": "TEST Firm"})
    assert r.status_code == 200, r.text
    s.csrf = r.json()["csrf_token"]
    s.email = email
    s.uid = r.json()["user_id"]
    return s


def cleanup_user(email):
    u = mdb.users.find_one({"email": email})
    if u:
        for col in ("clients", "payouts", "statements"):
            mdb[col].delete_many({"owner_id": u["user_id"]})
        mdb.payment_transactions.delete_many({"user_id": u["user_id"]})
        mdb.password_reset_tokens.delete_many({"user_id": u["user_id"]})
        mdb.user_sessions.delete_many({"user_id": u["user_id"]})
    mdb.users.delete_one({"email": email})


@pytest.fixture(scope="module")
def user():
    s = register()
    yield s
    cleanup_user(s.email)


@pytest.fixture(scope="module")
def user2():
    s = register()
    yield s
    cleanup_user(s.email)


@pytest.fixture(autouse=True, scope="module")
def clear_rate_limits():
    yield
    mdb.rate_limits.delete_many({})


# ---------- CORS ----------
def test_cors_evil_not_reflected():
    for r in (requests.options(f"{API}/auth/login", headers={"Origin": EVIL, "Access-Control-Request-Method": "POST"}),
              requests.get(f"{API}/health", headers={"Origin": EVIL})):
        assert r.headers.get("access-control-allow-origin") not in (EVIL, "*")


def test_cors_preview_allowed_with_csrf_header():
    r = requests.options(f"{API}/clients", headers={"Origin": BASE_URL, "Access-Control-Request-Method": "POST",
                                                    "Access-Control-Request-Headers": "content-type,x-csrf-token"})
    assert r.status_code == 200
    assert r.headers.get("access-control-allow-origin") == BASE_URL
    assert "x-csrf-token" in r.headers.get("access-control-allow-headers", "").lower()


# ---------- CSRF ----------
def test_csrf_blocks_without_token(user):
    assert user.post(f"{API}/clients", json={"name": "TEST_csrf"}).status_code == 403
    assert user.post(f"{API}/clients", json={"name": "TEST_csrf"}, headers={"X-CSRF-Token": "bad"}).status_code == 403
    assert user.put(f"{API}/clients/x", json={"name": "a"}).status_code == 403
    assert user.delete(f"{API}/clients/x").status_code == 403
    assert user.get(f"{API}/clients").status_code == 200  # GET not affected


def test_csrf_token_from_login_and_me(user):
    me = user.get(f"{API}/auth/me").json()
    assert me["csrf_token"] == user.csrf
    h = {"X-CSRF-Token": me["csrf_token"]}
    r = user.post(f"{API}/clients", json={"name": "TEST_csrf_ok"}, headers=h)
    assert r.status_code == 200
    cid = r.json()["id"]
    assert user.put(f"{API}/clients/{cid}", json={"name": "TEST_csrf_ok2"}, headers=h).status_code == 200
    assert user.delete(f"{API}/clients/{cid}", headers=h).status_code == 200
    # a different user's token must not work
    login = requests.post(f"{API}/auth/login", json=DEMO).json()["csrf_token"]
    assert user.post(f"{API}/clients", json={"name": "TEST_x"}, headers={"X-CSRF-Token": login}).status_code == 403


def test_csrf_no_cookie_returns_401_not_403():
    assert requests.post(f"{API}/clients", json={"name": "x"}).status_code == 401


def test_csrf_exempt_routes_with_cookies(user):
    assert user.post(f"{API}/auth/refresh").status_code == 200
    assert user.post(f"{API}/auth/forgot-password", json={"email": "nobody_test@example.com"}).status_code == 200
    assert user.post(f"{API}/auth/reset-password", json={"token": "nope", "password": "Abcdefgh1"}).status_code == 400
    assert user.post(f"{API}/webhooks/stripe", data=b"{}", headers={"stripe-signature": "bad"}).status_code == 400
    assert user.post(f"{API}/stripe/webhook", data=b"{}", headers={"stripe-signature": "bad"}).status_code == 400
    assert user.post(f"{API}/webhooks/textract", data=b"{}", headers={"X-LedgerStream-Signature": "bad"}).status_code == 401
    r = user.post(f"{API}/auth/login", json={"email": user.email, "password": "TestPass#2026"})
    assert r.status_code == 200
    user.csrf = r.json()["csrf_token"]  # new refresh cookie -> new csrf token


def test_logout_exempt():
    s = register()
    try:
        assert s.post(f"{API}/auth/logout").status_code == 200
    finally:
        cleanup_user(s.email)


# ---------- Forgot / reset ----------
def test_forgot_reset_flow():
    s = register()
    try:
        r = requests.post(f"{API}/auth/forgot-password", json={"email": s.email}, headers={"Origin": BASE_URL})
        assert r.status_code == 200 and r.json() == {"ok": True}
        rec = mdb.password_reset_tokens.find_one({"user_id": s.uid, "used": False})
        assert rec, "reset token not stored"
        token = rec["token"]
        logs = "".join(open(f, errors="ignore").read() for f in glob.glob("/var/log/supervisor/backend.*.log"))
        assert token not in logs, "reset token leaked into backend logs"
        assert requests.post(f"{API}/auth/reset-password", json={"token": "invalid", "password": "NewPass#2026"}).status_code == 400
        assert requests.post(f"{API}/auth/reset-password", json={"token": token, "password": "short"}).status_code == 422
        assert requests.post(f"{API}/auth/reset-password", json={"token": token, "password": "NewPass#2026"}).status_code == 200
        assert requests.post(f"{API}/auth/reset-password", json={"token": token, "password": "NewPass#2027"}).status_code == 400
        assert requests.post(f"{API}/auth/login", json={"email": s.email, "password": "TestPass#2026"}).status_code == 401
        assert requests.post(f"{API}/auth/login", json={"email": s.email, "password": "NewPass#2026"}).status_code == 200
        # unknown email: same response, no token
        r = requests.post(f"{API}/auth/forgot-password", json={"email": "nobody_x@example.com"})
        assert r.status_code == 200 and r.json() == {"ok": True}
    finally:
        mdb.login_attempts.delete_many({"identifier": {"$regex": s.email}})
        cleanup_user(s.email)
        mdb.rate_limits.delete_many({})


# ---------- Formula injection ----------
def test_qbo_and_variance_csv_neutralised(user):
    h = {"X-CSRF-Token": user.csrf}
    c = user.post(f"{API}/clients", json={"name": FORMULA, "platforms": ["toast"]}, headers=h).json()
    sample = user.get(f"{API}/payouts/sample/toast").content
    r = user.post(f"{API}/payouts", data={"client_id": c["id"], "platform": "toast", "commission_rate": "30"},
                  files={"file": ("=cmd.csv", io.BytesIO(sample), "text/csv")}, headers=h)
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    qbo = list(csv.reader(io.StringIO(user.get(f"{API}/payouts/{pid}/qbo").text)))
    cells = [cell for row in qbo[1:] for cell in row]
    assert any(cell == "'" + FORMULA for cell in cells), qbo[:3]
    assert not any(cell.startswith(("=", "+", "@")) for cell in cells), qbo[:3]
    var = list(csv.reader(io.StringIO(user.get(f"{API}/history/variance-export").text)))
    assert len(var) >= 2, var
    assert var[1][1] == "'" + FORMULA and var[1][3].startswith("'="), var[1]


def test_statement_export_neutralised(user):
    sid = f"stmt_test_{uuid.uuid4().hex[:8]}"
    mdb.statements.insert_one({"id": sid, "owner_id": user.uid, "client_id": "x", "filename": "TEST_s.pdf",
                               "status": "completed", "pages": 1, "rows": [], "created_at": "2026-01-01T00:00:00+00:00"})
    rows = [{"id": "r1", "date": "=1+1", "description": "@SUM(A1)", "reference": "+cmd", "withdrawal": 1.0},
            {"id": "r2", "date": "2026-01-02", "description": "-2+3", "reference": "\tX", "deposit": 2.0}]
    r = user.put(f"{API}/statements/{sid}/rows", json={"rows": rows}, headers={"X-CSRF-Token": user.csrf})
    assert r.status_code == 200, r.text
    out = list(csv.reader(io.StringIO(user.get(f"{API}/statements/{sid}/export?fmt=csv").text)))
    for row in out[1:]:
        for cell in row[:3]:
            assert not cell.startswith(("=", "+", "-", "@", "\t")), out
    x = user.get(f"{API}/statements/{sid}/export?fmt=xlsx")
    ws = load_workbook(io.BytesIO(x.content)).active
    for row in ws.iter_rows(min_row=2, values_only=True):
        for cell in row[:3]:
            assert not (isinstance(cell, str) and cell.startswith(("=", "+", "-", "@"))), row
    mdb.statements.delete_one({"id": sid})


# ---------- payments/status ----------
def test_payment_status_auth_and_ownership(user, user2):
    sid = f"cs_test_{uuid.uuid4().hex}"
    mdb.payment_transactions.insert_one({"session_id": sid, "user_id": user.uid, "plan": "starter",
                                         "status": "initiated", "payment_status": "pending"})
    try:
        assert requests.get(f"{API}/payments/status/{sid}").status_code == 401
        assert user2.get(f"{API}/payments/status/{sid}").status_code == 404
        r = user.get(f"{API}/payments/status/{sid}")
        assert r.status_code == 200 and r.json()["plan"] == "starter"
    finally:
        mdb.payment_transactions.delete_one({"session_id": sid})


# ---------- origin_url ----------
def test_checkout_origin(user):
    h = {"X-CSRF-Token": user.csrf}
    assert user.post(f"{API}/payments/checkout", json={"plan": "starter", "origin_url": EVIL}, headers=h).status_code == 400
    assert user.post(f"{API}/payments/checkout", json={"plan": "starter", "origin_url": "http://" + BASE_URL.split("://")[1]},
                     headers=h).status_code == 400
    r = user.post(f"{API}/payments/checkout", json={"plan": "starter", "origin_url": BASE_URL}, headers=h, timeout=60)
    assert r.status_code == 200 and r.json()["checkout_url"].startswith("https://checkout.stripe.com")


def test_checkout_origin_host_header_spoof(user):
    """X-Forwarded-Host / Host must not let an attacker whitelist their own domain."""
    h = {"X-CSRF-Token": user.csrf, "X-Forwarded-Host": "evil.example"}
    r = user.post(f"{API}/payments/checkout", json={"plan": "starter", "origin_url": EVIL}, headers=h, timeout=60)
    assert r.status_code == 400, f"spoofed X-Forwarded-Host accepted: {r.status_code}"


def test_portal_origin(user):
    mdb.users.update_one({"user_id": user.uid}, {"$set": {"plan": "starter", "stripe_customer_id": "cus_TESTfake"}})
    h = {"X-CSRF-Token": user.csrf}
    try:
        r = user.post(f"{API}/billing/portal", json={"origin_url": EVIL}, headers=h, timeout=60)
        assert r.status_code == 400 and "Unrecognised" in r.text, r.text
        r = user.post(f"{API}/billing/portal", json={"origin_url": BASE_URL}, headers=h, timeout=60)
        assert r.status_code != 400, r.text  # 502 expected for fake customer
    finally:
        mdb.users.update_one({"user_id": user.uid}, {"$set": {"plan": "free"}, "$unset": {"stripe_customer_id": ""}})


# ---------- Rate limits ----------
def test_forgot_rate_limit():
    mdb.rate_limits.delete_many({})
    codes = [requests.post(f"{API}/auth/forgot-password", json={"email": "nobody_rl@example.com"}).status_code for _ in range(6)]
    try:
        assert codes[:5] == [200] * 5 and codes[5] == 429, codes
    finally:
        mdb.rate_limits.delete_many({})


def test_forgot_rate_limit_xff_spoof():
    """Rotating X-Forwarded-For must not bypass the limit."""
    mdb.rate_limits.delete_many({})
    codes = [requests.post(f"{API}/auth/forgot-password", json={"email": "nobody_rl@example.com"},
                           headers={"X-Forwarded-For": f"10.9.{i}.1"}).status_code for i in range(7)]
    keys = mdb.rate_limits.distinct("key")
    mdb.rate_limits.delete_many({})
    assert 429 in codes, f"rate limit bypassed by spoofed XFF: {codes} keys={keys}"
